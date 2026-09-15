"""Read preserved historical exports; generate descriptive paper tables.

Run with the bundled Python runtime (openpyxl). No significance test is inferred
from summary statistics. Source annotations are retained only for explicit pairs.
"""
from pathlib import Path
import json, re, math, hashlib
import openpyxl

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'sources'
records, manifest = [], {}
values, evidence = {}, {}
for f in sorted(SRC.glob('*.xlsx')):
    rows = list(openpyxl.load_workbook(f, data_only=True).active.values)
    original = list(openpyxl.load_workbook(SRC / 'original_headers' / f.name, data_only=True).active.values)
    heads = rows[0]
    manifest[f.name] = {'sha256': hashlib.sha256(f.read_bytes()).hexdigest(),
                        'original_header_sha256': hashlib.sha256((SRC/'original_headers'/f.name).read_bytes()).hexdigest()}
    for ri, row in enumerate(rows[1:-1], 2):
        m = row[heads.index('M')]
        expected = 10 if '_10' in f.stem else 20
        for ci, h in enumerate(heads):
            match = re.fullmatch(r'([\deE+.\-]+) \(([\deE+.\-]+)\)(?: ([+=\-]))?', str(row[ci]))
            if not match:
                continue
            mean, sd, symbol = match.groups()
            # Renamed and original-header copies must have identical data cells.
            assert original[ri-1][ci] == row[ci], (f.name, ri, ci)
            rec = dict(file=f.name, cell=f'{openpyxl.utils.get_column_letter(ci+1)}{ri}',
                       problem=row[0], M=m, expected_M=expected, label=h,
                       implementation=original[0][ci], mean=mean, sd=sd,
                       symbol=symbol, comparator=heads[-1], included=m==expected)
            records.append(rec)
            if m != expected:
                continue
            key=(m,row[0],h)
            assert key not in values or values[key] == (mean,sd), key
            values[key]=(mean,sd)
            if symbol:
                pair=(m,row[0],h,heads[-1])
                assert pair not in evidence or evidence[pair]==symbol
                evidence[pair]=symbol

def names(m):
    k=15 if m==10 else 30
    return ['REMO', f'REMO (k={k})', 'REMO+HPC', f'REMO+HPC (k={k})',
            'REMO+candidate', f'full (k={k})'] + (['full'] if m==20 else [])
def problems(m):
    return [f'DTLZ{i}' for i in range(1,8)]+[f'WFG{i}' for i in range(1,10) if not (m==20 and i==4)]
def stats(m, before, after, subset=None):
    ps=subset or problems(m)
    ratios=[float(values[m,p,before][0])/float(values[m,p,after][0]) for p in ps]
    counts=[sum(x>1 for x in ratios),sum(x<1 for x in ratios),sum(x==1 for x in ratios)]
    syms=[]
    for p in ps:
        s=evidence.get((m,p,after,before))
        if s is None:
            old=evidence.get((m,p,before,after))
            s={'+':'-','-':'+','=':'='}.get(old)
        syms.append(s)
    return dict(means=counts, gm=math.exp(sum(map(math.log,ratios))/len(ps)),
                exported=[syms.count(s) for s in ['+','-','=']] if None not in syms else None)

summary={}
for m in [10,20]:
    ns=names(m)
    ids=['R6','RK','P6','PK','C6','FK']+(['F6'] if m==20 else [])
    lines=[r'\begin{table*}[p]',r'\centering',
           rf'\caption{{PAQC and CDIS component results for $M={m}$. Entries give mean IGD above standard deviation; the lowest mean in each row is bold. Configuration labels are defined in Table~\ref{{tab:exp:ablation-design}}.}}',
           rf'\label{{tab:exp:abl{m}}}',r'\footnotesize',r'\setlength{\tabcolsep}{5pt}',r'\renewcommand{\arraystretch}{1.12}',
           r'\begin{tabular}{l'+('c'*len(ns))+'}',r'\toprule','Problem & '+' & '.join(ids)+r' \\',r'\midrule']
    bests=[0]*len(ns)
    for p in problems(m):
        best=min(float(values[m,p,n][0]) for n in ns)
        cells=[]
        for i,n in enumerate(ns):
            a,b=values[m,p,n]
            if float(a)==best:
                bests[i]+=1
                a=r'\textbf{'+a+'}'
            cells.append(r'\shortstack{'+a+r'\\('+b+')}')
        lines.append(p+' & '+' & '.join(cells)+r' \\')
    lines += [r'\midrule','Lowest mean & '+' & '.join(map(str,bests))+r' \\',r'\bottomrule',r'\end{tabular}',r'\end{table*}']
    (ROOT/f'table_m{m}.tex').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    pairs=[('R6','P6'),('RK','PK'),('R6','C6'),('PK','FK'),('RK','FK'),('R6','RK'),('P6','PK')]
    if m==20: pairs += [('C6','F6'),('P6','F6'),('F6','FK')]
    mapping=dict(zip(ids,ns))
    summary[m]={}
    for a,b in pairs:
        s=stats(m,mapping[a],mapping[b])
        s['DTLZ']=stats(m,mapping[a],mapping[b],problems(m)[:7])['gm']
        s['WFG']=stats(m,mapping[a],mapping[b],problems(m)[7:])['gm']
        summary[m][a+'->'+b]=s

(ROOT/'source_snapshot.json').write_text(json.dumps({'manifest':manifest,'records':records},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(ROOT/'summary.json').write_text(json.dumps(summary,indent=2)+'\n',encoding='utf-8')
def triplet(x):
    return '/'.join(map(str,x)) if x is not None else '---'
lines=[r'\begin{table*}[t]',r'\centering',
       r'\caption{PAQC and CDIS component contrasts. Each arrow is evaluated from the second configuration against the first. Mean B/W counts problems with a lower/higher mean IGD. Export B/W/N counts reported better/worse/no-difference outcomes; a dash indicates that the direct comparison was not exported. These annotations have not been recomputed from run-level data. $G$ is defined in \eqref{eq:abl-ratio}.}',
       r'\label{tab:exp:abl-effects}',r'\small',r'\begin{tabular}{llcccccc}',r'\toprule',
       r'Component & Contrast & \multicolumn{3}{c}{$M=10$} & \multicolumn{3}{c}{$M=20$} \\',
       r'\cmidrule(lr){3-5}\cmidrule(lr){6-8}',r' & & Mean B/W & Export B/W/N & $G$ & Mean B/W & Export B/W/N & $G$ \\',r'\midrule']
for factor,pair in [('PAQC','R6->P6'),('PAQC','RK->PK'),('CDIS','R6->C6'),('CDIS','PK->FK'),('Both','RK->FK'),('PAQC','C6->F6'),('CDIS','P6->F6')]:
    a,b=pair.split('->')
    cells=[]
    for m in [10,20]:
        s=summary[m].get(pair)
        cells += [triplet(s['means'][:2]),triplet(s['exported']),f"{s['gm']:.3f}"] if s else ['---']*3
    lines.append(factor+' & '+a+r'$\to$'+b+' & '+' & '.join(cells)+r' \\')
lines += [r'\bottomrule',r'\end{tabular}',r'\end{table*}']
(ROOT/'table_effects.tex').write_text('\n'.join(lines)+'\n',encoding='utf-8')
lines=[r'\begin{table*}[t]',r'\centering',
       r'\caption{Effect of increasing the reported representative count from 6 to $K$. Mean B/W and Export B/W/N follow Table~\ref{tab:exp:abl-effects}. Ratios above one favor the larger count. DTLZ and WFG are summarized separately using \eqref{eq:abl-ratio}.}',
       r'\label{tab:exp:k-effects}',r'\small',r'\begin{tabular}{clccccc}',r'\toprule',
       r'$M$ & Contrast & Mean B/W & Export B/W/N & $G$ (all) & $G$ (DTLZ) & $G$ (WFG) \\',r'\midrule']
for m in [10,20]:
    for pair in ['R6->RK','P6->PK']+(['F6->FK'] if m==20 else []):
        s=summary[m][pair];a,b=pair.split('->')
        lines.append(str(m)+' & '+a+r'$\to$'+b+' & '+triplet(s['means'][:2])+' & '+triplet(s['exported'])+f" & {s['gm']:.3f} & {s['DTLZ']:.3f} & {s['WFG']:.3f}"+r' \\')
lines += [r'\bottomrule',r'\end{tabular}',r'\end{table*}']
(ROOT/'table_k.tex').write_text('\n'.join(lines)+'\n',encoding='utf-8')
print(json.dumps(summary,indent=2))
print('Source cells:',len(records),'unique included entries:',len(values))
