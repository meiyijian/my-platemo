"""Extract supplied IGD exports and build manuscript tables without editing Excel.

Usage: python build_tables.py --source-dir <directory containing the three xlsx>
       python build_tables.py  # rebuild from the committed, cell-addressed snapshot
Statistical symbols are supplied results, never inferred from means/stds.
"""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
OURS = "REMO_new2_AdaMaO_SDEOnly_UniformMix_Original"
ORDER = ["REMO", "PIEA", "CSEA", "PC-SAEA", "K-RVEA", "MCEA/D", "PACDIS"]
ALIASES = {"REMO": "REMO", "PIEA": "PIEA", "CSEA": "CSEA",
           "PCSAEA": "PC-SAEA", "PCSAEA_N100": "PC-SAEA",
           "KRVEA": "K-RVEA", "KRVEA_100": "K-RVEA",
           "MCEAD": "MCEA/D", OURS: "PACDIS"}
PROBLEMS = [f"DTLZ{i}" for i in range(1, 8)] + [f"WFG{i}" for i in range(1, 10)]
CELL = re.compile(r"^\s*([\d.]+e[+-]\d+)\s*\(([\d.]+e[+-]\d+)\)\s*([+=-])?\s*$", re.I)


def extract(source):
    from openpyxl import load_workbook
    records, manifest = [], {"excluded_algorithm": "R2AEA", "sources": []}
    for m in [10, 15, 20]:
        path = source / f"{m}目标.xlsx"
        content = path.read_bytes()
        wb = load_workbook(path, data_only=True)
        ws = wb["IGD"]
        headers = {str(c.value): c.column for c in ws[1] if c.value is not None}
        assert set(ALIASES[h] for h in headers if h in ALIASES) == set(ORDER)
        info = {"filename": path.name, "sha256": hashlib.sha256(content).hexdigest(),
                "sheet": "IGD", "M": m, "reference_algorithm": OURS,
                "headers": list(headers), "metadata": [], "exported_summary": {}}
        for row in range(2, 18):
            problem = ws.cell(row, headers["Problem"]).value
            assert problem == PROBLEMS[row - 2]
            assert ws.cell(row, headers["M"]).value == m
            meta = {k: ws.cell(row, headers[k]).value if k in headers else None
                    for k in ["N", "M", "D", "FE"]}
            info["metadata"].append({"problem": problem, **meta})
            for header, col in headers.items():
                if header not in ALIASES:
                    continue
                cell = ws.cell(row, col)
                match = CELL.fullmatch(str(cell.value))
                assert match, (path.name, cell.coordinate, cell.value)
                mean, std, symbol = match.groups()
                assert (symbol is None) == (header == OURS), (path.name, cell.coordinate)
                records.append({"M": m, "problem": problem, "algorithm": ALIASES[header],
                                "source_algorithm": header, "mean": mean, "std": std,
                                "symbol": symbol or "", "source_file": path.name,
                                "source_sheet": "IGD", "source_cell": cell.coordinate,
                                "raw": cell.value})
        for header, col in headers.items():
            if header in ALIASES and header != OURS:
                actual = Counter(r["symbol"] for r in records
                                 if r["M"] == m and r["source_algorithm"] == header)
                original = str(ws.cell(18, col).value)
                assert original == f"{actual['+']}/{actual['-']}/{actual['=']}", (m, header)
                info["exported_summary"][header] = original
        manifest["sources"].append(info)
        wb.close()
        assert path.read_bytes() == content, "Source workbook changed during extraction"
    with (HERE / "igd_snapshot.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=records[0].keys())
        writer.writeheader()
        writer.writerows(records)
    (HERE / "source_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path)
    args = parser.parse_args()
    if args.source_dir:
        extract(args.source_dir)
    with (HERE / "igd_snapshot.csv").open(encoding="utf-8-sig", newline="") as f:
        records = list(csv.DictReader(f))
    assert len(records) == 3 * 16 * 7
    summary = {"symbol_order": "+/-/=", "by_objectives": {}, "by_suite": {}}
    for m in [10, 15, 20]:
        data = [r for r in records if int(r["M"]) == m]
        lookup = {(r["problem"], r["algorithm"]): r for r in data}
        assert len(lookup) == 112
        totals = {}
        best = Counter()
        mean_comparisons = {}
        for algorithm in ORDER[:-1]:
            c = Counter(lookup[p, algorithm]["symbol"] for p in PROBLEMS)
            totals[algorithm] = [c["+"], c["-"], c["="]]
            mean_comparisons[algorithm] = sum(Decimal(lookup[p, "PACDIS"]["mean"]) < Decimal(lookup[p, algorithm]["mean"]) for p in PROBLEMS)
        for p in PROBLEMS:
            lowest = min(Decimal(lookup[p, a]["mean"]) for a in ORDER)
            best.update(a for a in ORDER if Decimal(lookup[p, a]["mean"]) == lowest)
        assert {r["source_algorithm"] for r in data if r["algorithm"] in ["PC-SAEA", "K-RVEA"]} == {"PCSAEA_N100", "KRVEA_100"}
        summary["by_objectives"][str(m)] = {"baseline_plus_minus_equal": totals,
            "best_mean_counts": dict(best), "lower_mean_counts": mean_comparisons,
            "best_mean_problems_PACDIS": [p for p in PROBLEMS if Decimal(lookup[p, "PACDIS"]["mean"]) == min(Decimal(lookup[p, a]["mean"]) for a in ORDER)]}
    lookup = {(r["problem"], int(r["M"]), r["algorithm"]): r for r in records}
    for suite, problems in [("DTLZ", PROBLEMS[:7]), ("WFG", PROBLEMS[7:])]:
        totals = {}
        for algorithm in ORDER[:-1]:
            counts = Counter(lookup[p, m, algorithm]["symbol"] for p in problems for m in [10, 15, 20])
            totals[algorithm] = [counts["+"], counts["-"], counts["="]]
        lines = ["% Generated from igd_snapshot.csv by build_tables.py.",
                 r"\begin{table*}[p]", r"\centering",
                 rf"\caption{{Comparison of IGD values on {suite}1--{len(problems)}.}}",
                 rf"\label{{tab:exp:{suite.lower()}}}", r"\footnotesize",
                 r"\setlength{\tabcolsep}{4pt}", r"\renewcommand{\arraystretch}{1.08}",
                 r"\begin{tabular}{@{}>{\centering\arraybackslash}m{32pt}>{\centering\arraybackslash}m{14pt}*{7}{>{\centering\arraybackslash}m{\dimexpr(\textwidth-110pt)/7\relax}}@{}}",
                 r"\toprule[1.2pt]", "Problem & $M$ & " + " & ".join(ORDER) + r" \\", r"\midrule"]
        for i, p in enumerate(problems):
            if i:
                lines.append(r"\midrule")
            for j, m in enumerate([10, 15, 20]):
                lowest = min(Decimal(lookup[p, m, a]["mean"]) for a in ORDER)
                cells = []
                for a in ORDER:
                    r = lookup[p, m, a]
                    sign = (r"\," + "$" + r["symbol"] + "$") if r["symbol"] else ""
                    line1, line2 = r["mean"], "(" + r["std"] + ")" + sign
                    prefix = ""
                    if Decimal(r["mean"]) == lowest:
                        prefix = r"\cellcolor{black!25}"
                        line1 = r"\textbf{" + line1 + "}"
                        line2 = r"\textbf{" + line2 + "}"
                    cells.append(prefix + r"\shortstack{" + line1 + r"\\" + line2 + "}")
                group = p if j == 1 else ""
                lines.append(group + " & " + str(m) + " & " + " & ".join(cells) + r" \\")
        lines += [r"\midrule", r"\multicolumn{2}{c}{$+/-/=$} & " + " & ".join("/".join(map(str, totals[a])) for a in ORDER[:-1]) + r" & --- \\",
                  r"\bottomrule[1.2pt]", r"\end{tabular}", r"\par\smallskip",
                  r"\begin{minipage}{\textwidth}\scriptsize",
                  r"Each entry shows the mean above the standard deviation in parentheses. Shaded bold cells have the lowest mean in their row. Symbols $+$, $-$ and $=$ indicate that the baseline is reported as better than, worse than or not significantly different from PACDIS, respectively. The last row totals these symbols in the stated order."]
        lines += [r"\end{minipage}", r"\end{table*}", ""]
        (HERE / f"table_{suite.lower()}.tex").write_text("\n".join(lines), encoding="utf-8")
        summary["by_suite"][suite] = {"baseline_plus_minus_equal": totals}
    (HERE / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
