"""Build the three-arm IGD+ ablation table from the Full-reference exports."""

from __future__ import annotations

import argparse
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path
from statistics import mean, stdev

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCES = {
    10: ROOT / "sources" / "ablation_full_reference_m10_igdp.xlsx",
    20: ROOT / "sources" / "ablation_full_reference_m20_igdp.xlsx",
}
OUTPUT = ROOT / "table_ablation_igdp.tex"
CN_MIRROR = ROOT.parents[1] / "HPDC-MaOEA_中文版.md"
CELL = re.compile(r"^(\d+\.\d+e[+-]\d+) \((\d+\.\d+e[+-]\d+)\)(?: ([+=-]))?$")
FAMILIES = {"DTLZ": range(1, 8), "WFG": range(1, 10)}
SOURCE_COLUMNS = {"Full": 8, "w/o CDIS": 6, "w/o PAQC": 7}
TABLE_COLUMNS = ("Full", "w/o CDIS", "w/o PAQC")


def read_audit(workbook) -> dict[tuple[str, str], list[float]]:
    """Check the 15-result sample sets against the workbook's audit sheet."""
    sheet = workbook["核对"]
    assert sheet.max_row == 66, sheet.max_row
    samples = {}
    for row in list(sheet.values)[2:]:
        problem, method = row[:2]
        name = "Full" if method == "full" else method
        if name not in SOURCE_COLUMNS:
            continue
        run_ids = [int(value) for value in row[3].split(",")]
        raw = [float(value) for value in row[10].split(",")]
        assert len(raw) == 20, (problem, name, len(raw))
        assert len(run_ids) == len(set(run_ids)) == 15, (problem, name, run_ids)
        assert all(1 <= run_id <= 20 for run_id in run_ids)
        selected = [raw[run_id - 1] for run_id in run_ids]
        if name == "Full":
            assert run_ids == list(range(1, 16)), (problem, run_ids)
        else:
            assert sorted(selected) == sorted(raw)[5:], (problem, name)
        # Run values and audit summaries are exported to four decimal places.
        assert abs(mean(selected) - row[4]) < 0.00011, (problem, name, "mean")
        assert abs(stdev(selected) - row[5]) < 0.00011, (problem, name, "sd")
        key = (problem, name)
        assert key not in samples, key
        samples[key] = selected
    assert len(samples) == 48, len(samples)
    return samples


def parse_cell(value: str, *, reference: bool = False) -> tuple[str, str, str | None]:
    match = CELL.fullmatch(value)
    if match is None:
        raise ValueError(f"Invalid workbook cell: {value!r}")
    mean, sd, symbol = match.groups()
    if (symbol is None) != reference:
        raise ValueError(f"Wrong reference/symbol status: {value!r}")
    return mean, sd, symbol


def read_sources() -> dict[tuple[str, int], dict]:
    records = {}
    for m, path in SOURCES.items():
        workbook = load_workbook(path, read_only=True, data_only=True)
        samples = read_audit(workbook)
        sheet = workbook["IGDp"]
        assert sheet.title == "IGDp", (path, sheet.title)
        rows = list(sheet.values)
        assert len(rows) == 18, (path, len(rows))
        assert rows[0] == (
            "Problem", "N", "M", "D", "FE", f"REMO (k={15 if m == 10 else 30})",
            "w/o CDIS", "w/o PAQC", "full",
        ), (path, rows[0])
        observed = {name: Counter() for name in TABLE_COLUMNS[1:]}
        for row in rows[1:-1]:
            problem, n, row_m, d, fe = row[:5]
            assert (n, row_m, fe) == (100, m, 300), (path, row[:5])
            assert d == (31 if problem in {"WFG2", "WFG3"} else 30), row[:5]
            key = (problem, m)
            assert key not in records, key
            cells = {
                name: parse_cell(row[index], reference=name == "Full")
                for name, index in SOURCE_COLUMNS.items()
            }
            reference_mean = Decimal(cells["Full"][0])
            for name, (mean, _, symbol) in cells.items():
                if symbol == "+":
                    assert Decimal(mean) < reference_mean, (key, name, symbol)
                elif symbol == "-":
                    assert Decimal(mean) > reference_mean, (key, name, symbol)
            for name, (_, _, symbol) in cells.items():
                if symbol is not None:
                    observed[name][symbol] += 1
            records[key] = {
                "D": d, "cells": cells,
                "samples": {name: samples[(problem, name)] for name in TABLE_COLUMNS},
            }
        for name, index in SOURCE_COLUMNS.items():
            if name == "Full":
                assert rows[-1][index] is None
                continue
            counts = "/".join(str(observed[name][s]) for s in "+-=")
            assert rows[-1][index] == counts, (path, name, rows[-1][index], counts)
        assert rows[-1][:5] == ("+/-/=", None, None, None, None)
        workbook.close()
    assert len(records) == 32, len(records)
    assert set(records) == {
        (f"{family}{i}", m)
        for family, ids in FAMILIES.items()
        for i in ids
        for m in SOURCES
    }
    return records


def tex_number(value: str) -> str:
    mantissa, exponent = value.split("e")
    return f"{mantissa}\\mathrm{{e}}{{{exponent}}}"


def tex_cell(cell: tuple[str, str, str | None], best: bool) -> str:
    mean, sd, symbol = cell
    mark = f"^{{{symbol}}}" if symbol is not None else ""
    prefix = "\\bestcell" if best else ""
    return f"{prefix}${tex_number(mean)}\\,({tex_number(sd)}){mark}$"


def make_table(records: dict) -> str:
    lines = [
        "\\begin{table*}[!t]",
        "\\centering",
        "\\caption{IGD$^{+}$ ablation results on DTLZ1--7 and WFG1--9 with 10 and 20 objectives "
        "($N=100$; $FE=300$; 15 runs per configuration).}",
        "\\label{tab:ablation_igdplus}",
        "\\small",
        "\\setlength{\\tabcolsep}{3pt}",
        "\\renewcommand{\\arraystretch}{1.35}",
        "\\begin{tabular*}{\\textwidth}{@{\\extracolsep{\\fill}}lccccc@{}}",
        "\\toprule",
        "Problem & $M$ & $D$ & Full & w/o CDIS & w/o PAQC \\\\",
        "\\midrule",
    ]
    counts = {name: Counter() for name in TABLE_COLUMNS[1:]}
    problems = [f"{family}{i}" for family, ids in FAMILIES.items() for i in ids]
    for index, problem in enumerate(problems):
        if problem == "WFG1":
            lines.append("\\midrule")
        elif index > 0:
            lines.append("\\addlinespace[2pt]")
        for m in SOURCES:
            entry = records[(problem, m)]
            cells = entry["cells"]
            minimum = min(Decimal(cells[name][0]) for name in TABLE_COLUMNS)
            parts = [problem if m == 10 else "", str(m), str(entry["D"])]
            for name in TABLE_COLUMNS:
                cell = cells[name]
                parts.append(tex_cell(cell, Decimal(cell[0]) == minimum))
                if cell[2] is not None:
                    counts[name][cell[2]] += 1
            lines.append(" & ".join(parts) + " \\\\")
    summary = ["\\multicolumn{2}{c}{$+/-/=$}", "--", "--"]
    summary.extend(
        f"${'/'.join(str(counts[name][symbol]) for symbol in '+-=')}$"
        for name in TABLE_COLUMNS[1:]
    )
    lines.extend([
        "\\midrule",
        " & ".join(summary) + " \\\\",
        "\\bottomrule",
        "\\end{tabular*}",
        "\\par\\smallskip",
        "\\begin{minipage}{\\textwidth}\\footnotesize",
        "Entries give mean (sample standard deviation) over 15 runs; lower IGD$^{+}$ is better. "
        "Light gray cells mark the lowest mean in each row, including displayed ties. "
        "$+$, $-$, and $=$ indicate that a control performs better than, worse than, "
        "or shows no detected difference from Full, respectively "
        "(two-sided unpaired rank-sum test, $p<0.05$; no documented multiple-comparison adjustment). "
        "Full denotes PACDIS and carries no comparison symbol. Full uses runs 1--15; "
        "each control uses the 15 runs remaining after excluding its five lowest "
        "IGD$^{+}$ values from 20 runs. Section~\\ref{sec:exp:ablation} defines both controls.",
        "\\end{minipage}",
        "\\end{table*}",
    ])
    return "\n".join(lines)


def make_markdown(records: dict) -> str:
    lines = [
        "**表 3　10 和 20 目标 DTLZ1–7 与 WFG1–9 上的 IGD$^+$ 消融结果"
        "（$N=100$；$FE=300$；每种配置统计 15 次结果）。**",
        "",
        "| 问题 | M | D | Full | w/o CDIS | w/o PAQC |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    counts = {name: Counter() for name in TABLE_COLUMNS[1:]}
    for family, ids in FAMILIES.items():
        for i in ids:
            problem = f"{family}{i}"
            for m in SOURCES:
                entry = records[(problem, m)]
                minimum = min(Decimal(entry["cells"][name][0]) for name in TABLE_COLUMNS)
                parts = [problem, str(m), str(entry["D"])]
                for name in TABLE_COLUMNS:
                    avg, sd, symbol = entry["cells"][name]
                    value = f"{avg} ({sd})"
                    if symbol is not None:
                        value += f"<sup>{symbol}</sup>"
                        counts[name][symbol] += 1
                    if Decimal(avg) == minimum:
                        value = f"**{value}**"
                    parts.append(value)
                lines.append("| " + " | ".join(parts) + " |")
    summary = ["+/-/=", "全部", "无", "无"]
    summary.extend("/".join(str(counts[name][s]) for s in "+-=") for name in TABLE_COLUMNS[1:])
    lines.extend([
        "| " + " | ".join(summary) + " |",
        "",
        "单元格给出 15 次结果的均值（样本标准差）；IGD$^+$ 越小越好。"
        "加粗对应原表浅灰单元格，表示各行最低均值，包括显示数值的并列情况。"
        "$+$、$-$、$=$ 分别表示对照优于 Full、劣于 Full 和未检测到差异"
        "（双侧非配对秩和检验，$p<0.05$；工作簿未记录多重比较校正）。"
        "Full 表示 PACDIS，不附比较符号。Full 采用 run 1–15；"
        "每个对照采用从 20 次运行中剔除其 IGD$^+$ 最低的 5 次后剩余的 15 次结果。"
        "第 4.3 节定义两种对照。",
    ])
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--sync-cn", action="store_true")
    args = parser.parse_args()
    records = read_sources()
    generated = (
        "% Generated from the two archived Full-reference IGD+ workbooks in sources/.\n"
        + make_table(records)
        + "\n"
    )
    if args.check:
        assert OUTPUT.read_text(encoding="utf-8") == generated
        assert make_markdown(records) in CN_MIRROR.read_text(encoding="utf-8")
        print("Ablation table matches both Full-reference workbooks (96 cells; 15 results per cell).")
    else:
        OUTPUT.write_text(generated, encoding="utf-8", newline="\n")
        print("Updated", OUTPUT)
    if args.sync_cn:
        assert not args.check, "Use --sync-cn to update, then --check to validate."
        text = CN_MIRROR.read_text(encoding="utf-8")
        start = text.index("**表 3　")
        end = text.index("### 4.4", start)
        CN_MIRROR.write_text(text[:start] + make_markdown(records) + "\n\n" + text[end:], encoding="utf-8", newline="\n")
        print("Synchronized Chinese ablation table.")
    for name in TABLE_COLUMNS[1:]:
        counts = Counter(
            records[key]["cells"][name][2] for key in records
        )
        print(name, "/".join(str(counts[s]) for s in "+-="))


if __name__ == "__main__":
    main()
