from pathlib import Path
import json, re, csv
import openpyxl

root=Path(r'C:\Users\lsx\Desktop\AdaMao实验表')
out=Path(__file__).parent
books={}
for p in sorted(root.rglob('*.xlsx')):
    if p.name.startswith('~$') or '_原始表格备份' in p.parts: continue
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    books[str(p.relative_to(root))]={s:[list(r) for r in wb[s].iter_rows(values_only=True)] for s in wb.sheetnames}
    wb.close()
(out/'books.json').write_text(json.dumps(books,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
for p,ss in books.items():
    print(p)
    for s,rows in ss.items():
        print(s,len(rows),'HEADER',rows[0], 'FIRST',rows[1], 'LAST',rows[-1])
tex=Path(r'D:\PlatEMO-master\论文写作\HPDC-MaOEA.tex')
backup=out/'HPDC-MaOEA.before.tex'
if not backup.exists(): backup.write_bytes(tex.read_bytes())
