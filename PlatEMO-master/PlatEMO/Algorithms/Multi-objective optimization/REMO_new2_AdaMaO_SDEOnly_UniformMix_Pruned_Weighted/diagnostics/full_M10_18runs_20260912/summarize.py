import json, re, csv
from pathlib import Path
import numpy as np
import openpyxl

out = Path(__file__).parent
runs = json.loads((out/'runs.json').read_text(encoding='utf-8'))
comparisons = json.loads((out/'comparisons.json').read_text())
problems = list(dict.fromkeys(r['problem'] for r in runs))
algs = list(dict.fromkeys(r['algorithm'] for r in runs if r['algorithm'] != 'Pruned'))
groups = {(a,p): [r for r in runs if r['algorithm']==a and r['problem']==p] for a in algs for p in problems}
assert all(len(rs)==18 and sorted(r['run'] for r in rs)==list(range(1,19)) for rs in groups.values())
assert all(r['IGD'] is not None and np.isfinite(r['IGD']) for r in runs)
means = np.array([[np.mean([r['IGD'] for r in groups[a,p]]) for a in algs] for p in problems])
assert all(len(set(row))==len(row) for row in means)
ranks = 1+np.argsort(np.argsort(means,axis=1),axis=1)
index = {(s['target'],s['comparator'],s['problem']):s for s in comparisons}
book = Path(r'C:\Users\lsx\Desktop\AdaMao实验表\参数简化版本\pruned_weight十目标.xlsx')
ws = openpyxl.load_workbook(book,data_only=True)['IGD']
discrepancies=[]
for pi,p in enumerate(problems):
 assert ws.cell(pi+2,1).value==p
 for ai,a in enumerate(algs):
  value=ws.cell(pi+2,ai+4).value
  match=re.match(r'([\d.eE+-]+) \(([\d.eE+-]+)\)(?: ([+=-]))?',value)
  displayed, sd, symbol=match.groups()
  calc=means[pi,ai]; calc_sd=np.std([r['IGD'] for r in groups[a,p]],ddof=1)
  if not np.isclose(float(displayed),calc,rtol=5.1e-5):discrepancies.append([p,a,'mean'])
  if not np.isclose(float(sd),calc_sd,rtol=.0051):discrepancies.append([p,a,'SD'])
  if a!='Weighted':
   s=index['Weighted',a,p]
   expected='=' if s['pDefault']>=.05 else ('-' if s['relativePercent']<0 else '+')
   if symbol!=expected:discrepancies.append([p,a,'symbol',symbol,expected])
assert not discrepancies, discrepancies

def wtl(target,other,key='pDefault'):
 ss=[index[target,other,p] for p in problems]
 return [sum(s[key]<.05 and s['relativePercent']<0 for s in ss),sum(s[key]>=.05 for s in ss),sum(s[key]<.05 and s['relativePercent']>0 for s in ss)]
def fmt_wtl(v):return '/'.join(map(str,v))
primary=[index['Weighted','Original',p] for p in problems]
stats={a:{'meanRank':float(ranks[:,i].mean()),'DTLZRank':float(ranks[:7,i].mean()),'WFGRank':float(ranks[7:,i].mean()),'bestCount':int((ranks[:,i]==1).sum())} for i,a in enumerate(algs)}
summary={'workbookReconciliation':{'means':144,'standardDeviations':144,'testSymbols':128,'mismatches':discrepancies},'rankings':stats,'weightedVsOriginal':{'meanBetter':sum(s['relativePercent']<0 for s in primary),'meanWorse':sum(s['relativePercent']>0 for s in primary),'WTL':wtl('Weighted','Original'),'HolmWTL':wtl('Weighted','Original','pHolm'),'medianPercent':float(np.median([s['relativePercent'] for s in primary]))},'baselineComparisons':{a:{'Weighted':wtl('Weighted',a),'Original':wtl('Original',a),'WeightedHolm':wtl('Weighted',a,'pHolm')} for a in algs[:7]},'actualFE':{a:sorted(set(r['FE'] for p in problems for r in groups[a,p])) for a in algs}}
(out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
with (out/'problem_mean_IGD_and_ranks.csv').open('w',newline='',encoding='utf-8-sig') as f:
 w=csv.writer(f);w.writerow(['problem']+[a+'_meanIGD' for a in algs]+[a+'_rank' for a in algs])
 for i,p in enumerate(problems):w.writerow([p]+means[i].tolist()+ranks[i].tolist())

lines=['# Pruned_Weighted 十目标 18 次实验分析','',
'核对日期：2026-09-12。范围：DTLZ1–7、WFG1–9，M=10，设置 N=100、maxFE=300；WFG2/3 实际 D=31，其余 D=30。',
'', '## 结论', '',
'Pruned_Weighted 保留了较强的整体竞争力，但目前不能认为它整体超过全参数版。相对于全参数版，16 题中均值改善与退化各 8 题；未经多重比较校正的胜/无显著差异/负为 3/10/3。对这 16 项精确秩和检验实施 Holm 校正后为 1/13/2：DTLZ7 改善，DTLZ2 和 DTLZ5 退化。',
'',
'九算法逐题均值排名的平均值：PIEA 3.3125，全参数版 3.6250，Weighted 3.8125。Weighted 为第三，全参数版第二；这是描述性排名，不能据此认定第二与第三存在显著总体差异。',
'', '## 数据与统计核对', '',
'- 九个主比较算法各使用相同 16 题、运行编号 1–18，共 2592 份 MAT；补充读取 Pruned 现有六题各 9 次，共 54 份。没有剔除离群运行。编号相同不能证明共同随机种子，因此逐题使用双侧非配对 Wilcoxon 秩和检验。',
'- Excel 的 144 个均值、144 个标准差以及 128 个符号均与 MAT 重算结果一致（按导出精度核对）。表格里的 “+” 表示该列对比算法优于最后一列 Weighted，“-” 表示 Weighted 更好。下文胜/平/负全部从被评价版本的角度统计；平只表示未检出差异。',
'- 表格复现使用 MATLAB ranksum 默认方法；同时保存 exact 方法及每个对手对应 16 题的 Holm 校正。本文校正结论限定在各自 16 题检验族，不表示控制所有算法、所有指标的总错误率。',
'- Weighted 和 Original 均为 288/288 次真实 FE=300。PIEA、MCEAD、PCSAEA_N100、KRVEA_100 也均为 300。REMO 为 300–310，CSEA 为 300–311，R2AEA 有 1 次 306、其余 300。因此与这些历史数据的比较不能称为所有运行严格等 FE；它们有额外评价机会。',
'- Weighted 未存 HV，故本次只得出 IGD 结论。使用保存的最终 IGD，不重新生成参考前沿；历史运行的源码哈希及全局随机种子未保存，当前设置文件不能替代每次运行的完整溯源。',
'', '## 与全参数版逐题比较', '',
'相对变化 = 100×(Weighted 均值 / Original 均值 − 1)，负值表示改善。95% 区间为各版本独立重采样 10000 次的均值比 percentile bootstrap 区间，没有多重校正。', '',
'| 问题 | Original IGD | Weighted IGD | 相对变化 | 精确 p | Holm p | 95% 相对变化区间 |',
'|---|---:|---:|---:|---:|---:|---:|']
for s in primary:
 lines.append(f"| {s['problem']} | {s['meanComparator']:.6g} | {s['meanTarget']:.6g} | {s['relativePercent']:+.2f}% | {s['pExact']:.5g} | {s['pHolm']:.5g} | [{s['ciLow']:+.2f}%, {s['ciHigh']:+.2f}%] |")
lines += ['',
'DTLZ7 的均值下降 27.43%，标准差由 2.798 降至 1.285，改善不仅表现为某次最优值。DTLZ2 与 DTLZ5 的均值分别上升 10.18% 和 25.57%，校正后仍有差异证据，不能用随机波动轻易解释。DTLZ4 上升 8.66%，未经校正 p=0.02236，但 Holm 后为 0.26831，应保留为退化信号。',
'',
'WFG 系列基本保留了原水平：两者在九算法中的 WFG 平均排名均为 3.2222。WFG1 改善 0.83%、WFG5 改善 4.69%，未通过 16 题 Holm 校正。不能把所有变化描述为显著改善，也不能把未显著差异当作等效证明。',
'', '## 与已有对比算法比较', '',
'下表采用与 Excel 一致的逐题 p<0.05，未校正。胜/平/负依次表示该版本显著优于对手、未检出差异、显著劣于对手。', '',
'| 对手 | Weighted 胜/平/负 | Original 胜/平/负 | Weighted 经逐对手 Holm 校正 |',
'|---|---:|---:|---:|']
for a in algs[:7]:lines.append(f"| {a} | {fmt_wtl(wtl('Weighted',a))} | {fmt_wtl(wtl('Original',a))} | {fmt_wtl(wtl('Weighted',a,'pHolm'))} |")
lines += ['',
'面对七个外部算法，共 112 项比较，Weighted 为 66/20/26，Original 为 68/22/22。该计数仅为描述性汇总，不把 112 项当作独立实验样本。Weighted 优于 REMO、MCEAD、CSEA、PCSAEA_N100、KRVEA_100 的问题多于劣于它们的问题，但对 PIEA 为 5/4/7。',
'', '| 算法 | 全部问题平均排名 | DTLZ 平均排名 | WFG 平均排名 | 均值第一的题数 |', '|---|---:|---:|---:|---:|']
for a in sorted(algs,key=lambda a:stats[a]['meanRank']):
 t=stats[a];lines.append(f"| {a} | {t['meanRank']:.4f} | {t['DTLZRank']:.4f} | {t['WFGRank']:.4f} | {t['bestCount']} |")
lines += ['',
'Weighted 在 WFG1 均值第一，DTLZ1、WFG5、WFG6、WFG7 均值第二。其明显短板包括 DTLZ3、DTLZ5、DTLZ6、WFG3：Weighted 与 PIEA 的 IGD 均值比约为 1.63、2.86、2.52、1.46。DTLZ7 虽较 Original 提升明显，仍落后于 KRVEA_100（1.6513）和 R2AEA（6.2038），自身为 8.7971。',
'', '## 关于恢复权重与参数简化的解释', '',
'仅恢复 0.75×关系质量+0.25×批次距离后，Weighted 对现有 Pruned 六题数据的均值在五题改善。但 Pruned 只有每题 9 次，不能把这一补充比较冒充全系列 18 对 18 的消融证据。六题检验中，DTLZ7 的改善最清楚；其它题的恢复仍不充分。',
'',
'Weighted 与 Original 不只相差是否保留辅助参数。当前总实验设置显示 Original 的 qKeep=0.80，Weighted=0.70；同时前者按 R+lambda×预测模糊度筛选并做批次加权，后者按 R 筛选并做批次加权。也就是说，筛选宽度和排序依据一起改变了。当前结果不能把 DTLZ2/5 的退化单独归因于 lambda0 删除，也不能得出参数删得越多效果必然越差。',
'',
'从源码可以提出待检验解释：qKeep 从 0.80 降到 0.70 通常把合格池从前约 20% 放宽到前约 30%；加权选择中的距离可能让较低关系得分的候选进入真实评价批次。恢复权重提供持续质量约束，但不等于恢复 Original 的候选排序与探索奖励。DTLZ2/5 的退化与这种解释相容，并不是该机制的因果证明。',
'',
'指标分支删去第二次 0.70 分位筛选，在合格候选足够多并且最终仍按同一指标取前六个时，不改变选择；最少 20 个和 nMin 的影响取决于它们实际是否激活。没有分支日志，不应把它们断言为零影响，也不应优先凭空恢复。',
'', '## 下一步建议', '',
'1. 保留当前 Weighted 为固定的简化候选，保留 Original 作为完整性能参考。现有 18 次足以揭示主要取舍，暂不需要为了让均值好看而反复追加这 16 题。',
'2. 优先隔离 qKeep：在 Weighted 上只比较 0.70 与 0.80，先覆盖 DTLZ2、DTLZ5、DTLZ7、WFG7，记录共同初始设计/种子、真实 FE、批次质量与分散性，重复次数预先固定。原版本作为第三个参照；qKeep=0.80 的 Weighted 仍不等于 Original。',
'3. 若收紧 qKeep 仍不能恢复 DTLZ2/5，再单独恢复模糊度奖励及其门控，避免一次加回多项规则。',
'4. 若目标是验证高目标泛化，冻结后扩展到 M=15/20，至少同时比较 Weighted 与 Original；此前六题参与过版本选择，后续新问题/新目标数的确认更有价值。主比较应尽量统一实际 FE 并补齐 HV。',
'', '## 文件', '',
f'- 汇总表：`{book}`，IGD!A1:L18。',
'- 原始数据：`C:\\Users\\lsx\\Desktop\\REMOandDREMO测试集\\10目标\\n30`。',
'- 本目录 `run_metrics.csv` 为逐次指标及实际 FE，`comparisons.csv` 为逐题统计，`problem_mean_IGD_and_ranks.csv` 为均值与排名。',
'- 只读分析正式结果与算法源文件，未修改原始 MAT、Excel 或算法。', '']
(out/'Pruned_Weighted_十目标18次_分析报告.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
