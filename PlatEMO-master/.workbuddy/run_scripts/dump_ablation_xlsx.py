import sys, glob, os
import openpyxl

for path in sys.argv[1:]:
    print("=" * 100)
    print("FILE:", path)
    if not os.path.isfile(path):
        print("  MISSING")
        continue
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print("  SHEET:", ws.title, ws.max_row, "x", ws.max_column)
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
            cells = []
            for c in r:
                v = c.value
                cells.append("" if v is None else str(v))
            print("   | " + " | ".join(cells))
