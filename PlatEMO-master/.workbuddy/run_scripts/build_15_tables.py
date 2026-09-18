# -*- coding: utf-8 -*-
"""Build the 15-objective IGD and IGDp tables in the format of the reference
workbook lambdat030nobatch十目标IGDp.xlsx.

Reads the formatted cell texts produced by BuildTable15Data.m and reproduces the
conventions of the reference sheet: Times New Roman 11, centred, thin borders
over the whole block, column widths and row height taken from the reference
file, the label "+/-/=" merged over A:C of the last row, and the per-row minimum
painted in blue (#3333E9). Eight algorithm columns, the NoBatchDist variant last.
"""

import csv
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

REF_FILE = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本\lambdat030nobatch十目标IGDp.xlsx"
OUT_DIR = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本"
CSV_TMPL = os.path.join(OUT_DIR, "table15_{metric}.csv")

ALGS = ["REMO", "PIEA", "MCEAD", "CSEA", "PCSAEA_N100", "KRVEA_100",
        "REMO_UniformMix_Pruned_Weighted_Lambdat030",
        "REMO_UniformMix_Pruned_Weighted_Lambdat030_NoBatchDist"]
LABELS = ["REMO", "PIEA", "MCEAD", "CSEA", "PCSAEA_N100", "KRVEA_100",
          "Lambdat030", "NoBatchDist"]
BLUE = "3333E9"
PROBN = 16
SUMMARY_ROW = PROBN + 2          # row 18
FIRST_ALG_COL = 4                # A Problem, B M, C D, D.. algorithms

JOBS = [
    ("IGD",  "lambdat030nobatch十五目标.xlsx",     "IGD"),
    ("IGDp", "lambdat030nobatch十五目标IGDp.xlsx", "IGDp"),
]


def reference_format():
    wb = openpyxl.load_workbook(REF_FILE)
    ws = wb[wb.sheetnames[0]]
    widths = {}
    for key, dim in ws.column_dimensions.items():
        if dim.width:
            widths[key] = dim.width
    height = ws.row_dimensions[1].height or 20
    f = ws["A1"].font
    return widths, height, (f.name or "Times New Roman", f.size or 11)


def build(metric, out_name, sheet_name):
    csv_path = CSV_TMPL.format(metric=metric)
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if len(rows) != PROBN:
        raise SystemExit(f"expected {PROBN} rows in {csv_path}, found {len(rows)}")

    widths, height, (fname, fsize) = reference_format()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center")

    def style(cell, blue=False):
        cell.font = Font(name=fname, size=fsize, color=(BLUE if blue else "000000"))
        cell.alignment = centre
        cell.border = border

    for j, text in enumerate(["Problem", "M", "D"] + ALGS, start=1):
        style(ws.cell(row=1, column=j, value=text))

    for i, r in enumerate(rows):
        excel_row = i + 2
        style(ws.cell(row=excel_row, column=1, value=r["Problem"]))
        style(ws.cell(row=excel_row, column=2, value=int(float(r["M"]))))
        style(ws.cell(row=excel_row, column=3, value=int(float(r["D"]))))
        best = r["BestAlg"].strip()
        for k, alg in enumerate(ALGS):
            style(ws.cell(row=excel_row, column=FIRST_ALG_COL + k, value=r[alg]),
                  blue=(LABELS[k] == best))

    style(ws.cell(row=SUMMARY_ROW, column=1, value="+/-/="))
    for j in (2, 3):
        style(ws.cell(row=SUMMARY_ROW, column=j, value=None))
    ws.merge_cells(start_row=SUMMARY_ROW, start_column=1,
                   end_row=SUMMARY_ROW, end_column=3)
    for k, lab in enumerate(LABELS[:-1]):
        style(ws.cell(row=SUMMARY_ROW, column=FIRST_ALG_COL + k,
                      value=rows[0][f"cnt_{lab}"]))
    style(ws.cell(row=SUMMARY_ROW, column=FIRST_ALG_COL + len(ALGS) - 1, value=None))

    alg_cols = [get_column_letter(FIRST_ALG_COL + k) for k in range(len(ALGS))]
    for key, width in widths.items():
        if key in "ABC" or key in alg_cols:
            ws.column_dimensions[key].width = width
    for excel_row in range(1, SUMMARY_ROW + 1):
        ws.row_dimensions[excel_row].height = height

    out_path = os.path.join(OUT_DIR, out_name)
    wb.save(out_path)
    print(f"{metric}: wrote {out_path} (sheet '{sheet_name}', {len(rows)} problems, "
          f"{len(ALGS)} algorithms)")


if __name__ == "__main__":
    for metric, out_name, sheet_name in JOBS:
        build(metric, out_name, sheet_name)
    print("XLSX BUILD DONE")
