# -*- coding: utf-8 -*-
"""Dump the header and a few rows of the three source workbooks of the paper table."""
import os

import openpyxl

BASE = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本"
FILES = {10: "lambdat030nobatch十目标IGDp.xlsx",
         15: "lambdat030nobatch十五目标IGDp.xlsx",
         20: "lambdat030nobatch二十目标IGDp.xlsx"}

for m, name in FILES.items():
    p = os.path.join(BASE, name)
    print("=" * 120)
    print("M=%d  %s  exists=%s" % (m, name, os.path.exists(p)))
    if not os.path.exists(p):
        continue
    wb = openpyxl.load_workbook(p)
    print("  sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        print("  --- sheet %s  dims %s  merged %s" % (ws.title, ws.dimensions,
                                                      [str(x) for x in ws.merged_cells.ranges]))
        for r in range(1, min(ws.max_row, 4) + 1):
            vals = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                vals.append("" if v is None else str(v))
            print("     r%-3d %s" % (r, " | ".join(vals)))
        print("     ... last row r%d: %s" % (ws.max_row,
              " | ".join("" if ws.cell(row=ws.max_row, column=c).value is None
                         else str(ws.cell(row=ws.max_row, column=c).value)
                         for c in range(1, ws.max_column + 1))))
        print("     col widths:", {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width})
        f = ws["A1"].font
        print("     A1 font:", f.name, f.sz, "| freeze:", ws.freeze_panes)
