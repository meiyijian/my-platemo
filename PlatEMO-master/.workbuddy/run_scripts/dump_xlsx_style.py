import sys
import openpyxl
from openpyxl.utils import get_column_letter

for path in sys.argv[1:]:
    print("=" * 100)
    print("FILE:", path)
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        print("  SHEET:", ws.title, "dims", ws.dimensions, "freeze", ws.freeze_panes)
        print("  merged:", [str(m) for m in ws.merged_cells.ranges])
        print("  col widths:", {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width})
        print("  row heights:", {k: v.height for k, v in ws.row_dimensions.items() if v.height})
        for r in range(1, min(ws.max_row, 4) + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is None and cell.border.left.style is None:
                    continue
                f = cell.font
                al = cell.alignment
                print("   %s%d val=%r font(name=%s,sz=%s,b=%s,color=%s) align(h=%s,v=%s,wrap=%s) numfmt=%s border(l=%s,r=%s,t=%s,b=%s)" % (
                    get_column_letter(c), r, cell.value,
                    f.name, f.sz, f.b, f.color.rgb if f.color else None,
                    al.horizontal, al.vertical, al.wrap_text, cell.number_format,
                    cell.border.left.style, cell.border.right.style,
                    cell.border.top.style, cell.border.bottom.style))
            print("   ---")
        # last rows
        for r in range(max(1, ws.max_row - 1), ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                f = cell.font
                print("   LAST %s%d val=%r font(name=%s,sz=%s,color=%s) align=%s/%s border=%s" % (
                    get_column_letter(c), r, cell.value,
                    f.name, f.sz, f.color.rgb if f.color else None,
                    cell.alignment.horizontal, cell.alignment.vertical,
                    cell.border.left.style))
