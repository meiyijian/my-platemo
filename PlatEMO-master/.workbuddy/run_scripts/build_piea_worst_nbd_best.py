# -*- coding: utf-8 -*-
"""Build the "PIEA worst vs NoBatchDist" IGD+ comparison tables, two flavours.

Flavour A (agg='min')  : last column = NoBatchDist's BEST run  (20 runs)
Flavour B (agg='mean') : last column = NoBatchDist's MEAN over 20 runs

Shared rules (fixed by the user, 2026-09-20):
  * the last column of the reference workbook is replaced by
    REMO_UniformMix_Pruned_Weighted_Lambdat030_NoBatchDist;
  * the PIEA column is replaced by the WORST (maximum) IGDp over its 30 runs;
  * every other algorithm keeps the value already printed in the reference
    workbook, verbatim;
  * the bracketed standard deviation is kept verbatim from the reference
    workbook (PIEA's own std for the PIEA column, PACDIS's std for the last
    column);
  * the +/-/= symbol of every baseline is recomputed: Mann-Whitney U
    (MATLAB ranksum equivalent, asymptotic + tie correction) of that
    baseline's runs against NoBatchDist's 20 runs, truncated to the common
    length; p >= 0.05 -> '=', otherwise the direction follows the two numbers
    actually printed in the table (min-is-better);
  * the per-row minimum of the seven printed numbers is painted blue.
"""
import glob
import os
import re
from copy import copy

import numpy as np
import openpyxl
import scipy.io as sio
from scipy.stats import mannwhitneyu

REF_DIR = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本"
OUT_DIR = os.path.join(REF_DIR, "PIEA最差版本与Nobatch最好版本")
TABDIR = r"C:\Users\lsx\Desktop\AdaMao实验表"
RAWDIR = r"C:\Users\lsx\Desktop\REMOandDREMO测试集"

PROBS = ["DTLZ1", "DTLZ2", "DTLZ3", "DTLZ4", "DTLZ5", "DTLZ6", "DTLZ7",
         "WFG1", "WFG2", "WFG3", "WFG4", "WFG5", "WFG6", "WFG7", "WFG8", "WFG9"]
BASE_ALGS = ["REMO", "PIEA", "MCEAD", "CSEA", "PCSAEA_N100", "KRVEA_100"]
NBD = "REMO_UniformMix_Pruned_Weighted_Lambdat030_NoBatchDist"
LAST_ALG = "REMO_UniformMix_Pruned_Weighted_Lambdat030"
BLUE = "003333E9"
P_THRESHOLD = 0.05

# flavour -> (aggregator of the 20 NoBatchDist runs, filename suffix, log label)
FLAVOURS = {
    "min": (np.nanmin, "最好", "best "),
    "mean": (np.nanmean, "均值", "mean "),
}

CELL_RE = re.compile(r"^([\d.]+e[+-]\d+)\s*\(([\d.]+e[+-]\d+)\)(?:\s*([+\-=]))?$")


def fmt(value):
    """MATLAB sprintf('%.4e') with the 'e+0'/'e-0' folding used by the reference table."""
    return ("%.4e" % value).replace("e-0", "e-").replace("e+0", "e+")


def parse_cell(txt):
    """-> (value, std_text, sign, 'value (std)' purity preserved)."""
    m = CELL_RE.match(txt.strip())
    if not m:
        raise ValueError("unparsable cell: %r" % txt)
    return float(m.group(1)), m.group(2), (m.group(3) or ""), "%s (%s)" % (m.group(1), m.group(2))


def raw_nbd_dir(m):
    if m == 10:
        return os.path.join(RAWDIR, "10目标", "n30", NBD)
    return os.path.join(RAWDIR, "%d目标" % m, NBD)


def nbd_runs(m, prob):
    """The per-run final IGDp of NoBatchDist, in run order."""
    rd = raw_nbd_dir(m)
    out = []
    for r in range(1, 31):
        hits = sorted(glob.glob(os.path.join(rd, "%s_%s_M%d_D*_%d.mat" % (NBD, prob, m, r))))
        if not hits:
            continue
        series = np.ravel(np.asarray(
            sio.loadmat(hits[0], variable_names=["metric"])["metric"]["IGDp"][0, 0],
            dtype=float))
        out.append((r, float(series[-1])))
    return out


def baseline_runs(m, alg, prob):
    f = os.path.join(TABDIR, "IGDplus_M%d" % m, alg, "%s_IGDp_M%d.mat" % (prob, m))
    v = np.ravel(sio.loadmat(f, variable_names=["IGDpFinal"])["IGDpFinal"]).astype(float)
    return v[~np.isnan(v)]


def build(m, agg="min"):
    aggregator, suffix, label = FLAVOURS[agg]

    ref_path = os.path.join(REF_DIR, "IGDp%d目标.xlsx" % m)
    wb = openpyxl.load_workbook(ref_path)
    ws = wb["IGDp"]

    header = [ws.cell(row=1, column=c).value for c in range(1, 11)]
    assert header[3:9] == BASE_ALGS, header
    assert header[9] == LAST_ALG, header

    cnt = {a: [0, 0, 0] for a in BASE_ALGS}   # + / - / =
    log = []

    for i, prob in enumerate(PROBS):
        row = i + 2
        old_best_val, old_best_std, _, _ = parse_cell(ws.cell(row=row, column=10).value)

        nbd = nbd_runs(m, prob)
        nbd_vals = np.array([v for _, v in nbd], dtype=float)
        nbd_disp = float(aggregator(nbd_vals))
        nbd_info = "run %d/%d" % (nbd[int(np.nanargmin(nbd_vals))][0], len(nbd_vals))

        runs_cache = {a: baseline_runs(m, a, prob) for a in BASE_ALGS}

        shown = {}       # column -> (display value, text to write)
        for k, alg in enumerate(BASE_ALGS):
            col = 4 + k
            val_old, std_old, _, pure = parse_cell(ws.cell(row=row, column=col).value)
            if alg == "PIEA":
                runs = runs_cache[alg]
                disp = float(np.nanmax(runs))
                argrun = int(np.nanargmax(runs)) + 1
                shown[col] = (disp, "%s (%s)" % (fmt(disp), std_old))
                log.append("    %-7s PIEA worst=%-13s (run %d/%d, ref mean was %s)"
                           % (prob, fmt(disp), argrun, len(runs), fmt(val_old)))
            else:
                shown[col] = (val_old, pure)     # untouched, symbol re-attached below

        shown[10] = (nbd_disp, "%s (%s)" % (fmt(nbd_disp), old_best_std))
        log.append("    %-7s NBD  %s=%-13s (%s, own std=%.2e, ref PACDIS mean was %s, %s)"
                   % (prob, label.strip().ljust(4), fmt(nbd_disp), nbd_info,
                      float(np.nanstd(nbd_vals)), fmt(old_best_val), old_best_std))

        minlen = min([len(nbd_vals)] + [len(v) for v in runs_cache.values()])
        for k, alg in enumerate(BASE_ALGS):
            col = 4 + k
            x = runs_cache[alg][:minlen]
            y = nbd_vals[:minlen]
            pval = float(mannwhitneyu(x, y, alternative="two-sided",
                                      method="asymptotic").pvalue)
            disp = shown[col][0]
            if pval >= P_THRESHOLD or disp == nbd_disp:
                sign = "="
            elif disp < nbd_disp:      # IGD+ is min-is-better
                sign = "+"
            else:
                sign = "-"
            cnt[alg][{"+": 0, "-": 1, "=": 2}[sign]] += 1
            shown[col] = (disp, "%s %s" % (shown[col][1], sign))
            log.append("      %-7s p=%.4f -> %s   (%s  vs  %s)"
                       % (alg, pval, sign, fmt(disp), fmt(nbd_disp)))

        best_val = min(v for v, _ in shown.values())
        for col, (disp, txt) in sorted(shown.items()):
            cell = ws.cell(row=row, column=col, value=txt)
            base = copy(cell.font)
            base.color = openpyxl.styles.Color(rgb=(BLUE if disp == best_val else "000000"))
            cell.font = base

    ws.cell(row=1, column=10, value=NBD)
    for k, alg in enumerate(BASE_ALGS):
        ws.cell(row=18, column=4 + k, value="%d/%d/%d" % tuple(cnt[alg]))
    ws.cell(row=18, column=10, value=None)

    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, "IGDp%d目标_PIEA最差vsNoBatchDist%s.xlsx" % (m, suffix))
    wb.save(out)

    print("=" * 110)
    print("M = %d  [NoBatchDist = %s]  ->  %s" % (m, agg, out))
    print("  +/-/= counters: %s" % {a: "%d/%d/%d" % tuple(cnt[a]) for a in BASE_ALGS})
    print("\n".join(log))


if __name__ == "__main__":
    for agg in ("min", "mean"):
        for m in (10, 15, 20):
            build(m, agg)
    print("\nALL DONE")
