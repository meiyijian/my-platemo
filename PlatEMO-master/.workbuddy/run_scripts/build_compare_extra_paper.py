import json
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

JSON_DIR = r"D:\PlatEMO-master\PlatEMO-master\.workbuddy\ablation_logs"
OUT_DIR = r"C:\Users\lsx\Desktop\AdaMao实验表\SSDE对比"
CN = {10: "十目标", 15: "十五目标", 20: "二十目标"}

THIN = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="Times New Roman", size=11)
FONT_BEST = Font(name="Times New Roman", size=11, color="FF3333E9")
ALIGN = Alignment(horizontal="center", vertical="center")
WIDTHS = {"A": 10.62, "B": 6.62}
ALG_WIDTH = 22.62


def build_one(payload, out_path):
    M = payload["M"]
    problems, Ds = payload["problems"], payload["Ds"]
    t = payload["tables"]
    labels, counts, best = t["labels"], t["counts"], t["best"]
    cells = [row.split("|") for row in t["cells"]]
    nA, nP = len(labels), len(problems)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IGDp"
    for c, txt in enumerate(["Problem", "N", "M", "D", "FE"] + labels, start=1):
        ws.cell(row=1, column=c, value=txt)
    for p in range(nP):
        r = 2 + p
        ws.cell(row=r, column=1, value=problems[p])
        ws.cell(row=r, column=2, value=100)
        ws.cell(row=r, column=3, value=M)
        ws.cell(row=r, column=4, value=int(Ds[p]))
        ws.cell(row=r, column=5, value=300)
        for a in range(nA):
            ws.cell(row=r, column=6 + a, value=cells[p][a])
        ws.cell(row=r, column=6 + (int(best[p]) - 1)).font = FONT_BEST

    last = 2 + nP
    ws.cell(row=last, column=1, value="+/-/=")
    ws.merge_cells(start_row=last, start_column=1, end_row=last, end_column=5)
    for a in range(nA):
        ws.cell(row=last, column=6 + a, value=counts[a] if a < len(counts) else "")

    for r in range(1, last + 1):
        ws.row_dimensions[r].height = 15.0
        for c in range(1, 6 + nA):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = ALIGN
            if cell.font.color is None or cell.font.color.rgb != "FF3333E9":
                cell.font = FONT
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    for a in range(nA):
        ws.column_dimensions[get_column_letter(6 + a)].width = ALG_WIDTH
    wb.save(out_path)
    return {"file": out_path, "labels": labels, "counts": counts}


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    for M in (10, 15, 20):
        jf = os.path.join(JSON_DIR, "compare_SSDE_M%d.json" % M)
        if not os.path.isfile(jf):
            print("MISSING %s" % jf)
            continue
        payload = json.load(open(jf, encoding="utf-8"))
        out = os.path.join(OUT_DIR, "SSDE_vs论文七算法_%sIGDp.xlsx" % CN[M])
        info = build_one(payload, out)
        print("wrote %s" % info["file"])
        print("   labels : %s" % " | ".join(info["labels"]))
        print("   counts : %s" % " | ".join(info["counts"]))


if __name__ == "__main__":
    main()
