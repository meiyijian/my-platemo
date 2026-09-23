# -*- coding: utf-8 -*-
"""Build the three FE500 IGDp comparison tables as formatted xlsx.

One workbook per objective count (M=10/15/20). Each row is the mean (std) of the
last-snapshot IGD+ over runs 1-10, with the +/-/= symbol of a MATLAB ranksum test
against the last column (the NoBatchDist variant, displayed as PACDIS).
The lowest mean of each row is written in blue, the summary row carries the
per-column +/-/= counts, and the PACDIS column has no symbol by construction.
"""
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, Side

SRC = r"D:\PlatEMO-master\PlatEMO-master\.workbuddy\run_scripts"
DST = r"C:\Users\lsx\Desktop\AdaMao实验表\nobatchdict版本"
JOBS = [
    (10, "fe500_m10_igdp_table.csv", "nobatchdict以PACDIS为基准十目标IGDp.xlsx"),
    (15, "fe500_m15_igdp_table.csv", "nobatchdict以PACDIS为基准十五目标IGDp.xlsx"),
    (20, "fe500_m20_igdp_table.csv", "nobatchdict以PACDIS为基准二十目标IGDp.xlsx"),
]
ALG_COLS = ["REMO", "SSDE", "PCSAEA", "SAMOEATL2M", "CSEA", "HES_EA",
            "REMO_UniformMix_Pruned_Weighted_Lambdat030_NoBatchDist"]
HEADERS = ["REMO", "SSDE", "PC-SAEA", "SAMOEA-TL2M", "CSEA", "HES-EA", "PACDIS"]
CNT_COLS = ["cnt_REMO", "cnt_SSDE", "cnt_PC-SAEA", "cnt_SAMOEA-TL2M",
            "cnt_CSEA", "cnt_HES-EA"]
LEAD = ["Problem", "M", "D"]

BLUE = Color(rgb="FF3333E9")
BLACK = Color(rgb="FF000000")
MEAN = re.compile(r"^\s*([\d.]+e[+-]\d+)")
_thin = Side(style="thin", color="FF000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center")


def build(m, csv_name, out_name):
    with open(os.path.join(SRC, csv_name), encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 16, (csv_name, len(rows))

    wb = Workbook()
    ws = wb.active
    ws.title = "IGDp"

    header = LEAD + HEADERS
    for c, text in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=text)
        cell.font = Font(name="Times New Roman", size=11, bold=True, color=BLACK)
        cell.alignment = CENTER
        cell.border = BORDER

    for r, row in enumerate(rows, start=2):
        assert int(row["M"]) == m
        ws.cell(row=r, column=1, value=row["Problem"])
        ws.cell(row=r, column=2, value=int(row["M"]))
        ws.cell(row=r, column=3, value=int(row["D"]))
        means = {}
        for c, col in enumerate(ALG_COLS, start=4):
            value = row[col]
            ws.cell(row=r, column=c, value=value)
            hit = MEAN.match(value)
            if hit:
                means[c] = float(hit.group(1))
        lowest = min(means.values())
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=r, column=c)
            blue = means.get(c) == lowest
            cell.font = Font(name="Times New Roman", size=11,
                             color=BLUE if blue else BLACK)
            cell.alignment = CENTER
            cell.border = BORDER

    summary = 18
    ws.cell(row=summary, column=1, value="+/-/=")
    for c, col in enumerate(CNT_COLS, start=4):
        ws.cell(row=summary, column=c, value=rows[0][col])
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=summary, column=c)
        cell.font = Font(name="Times New Roman", size=11, color=BLACK)
        cell.alignment = CENTER
        cell.border = BORDER

    for c, width in enumerate([10.6, 6.6, 6.6] + [24.0] * 6 + [26.0], start=1):
        ws.column_dimensions[chr(64 + c)].width = width

    out = os.path.join(DST, out_name)
    wb.save(out)
    counts = " | ".join("%s %s" % (h, rows[0][k]) for h, k in zip(HEADERS[:-1], CNT_COLS))
    print("wrote %s" % out)
    print("   M=%d  rows=%d  +/-/=: %s" % (m, len(rows), counts))


def main():
    if not os.path.isdir(DST):
        os.makedirs(DST)
        print("created %s" % DST)
    for m, csv_name, out_name in JOBS:
        build(m, csv_name, out_name)


if __name__ == "__main__":
    main()
