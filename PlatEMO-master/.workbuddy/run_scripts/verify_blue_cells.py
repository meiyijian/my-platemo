import sys, os
import openpyxl
from openpyxl.utils import get_column_letter

for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    print("=" * 90)
    print(os.path.basename(path))
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print("  header:", hdr)
    bad = 0
    for r in range(2, 18):
        vals = {}
        for c in range(6, ws.max_column + 1):
            txt = ws.cell(row=r, column=c).value
            vals[c] = txt
        # blue cell
        blue = [c for c in vals if ws.cell(row=r, column=c).font.color is not None
                and ws.cell(row=r, column=c).font.color.rgb == "FF3333E9"]
        # numeric mins
        nums = {c: float(str(t).split()[0]) for c, t in vals.items()}
        nm = min(nums, key=nums.get)
        ok = (len(blue) == 1 and blue[0] == nm)
        if not ok:
            bad += 1
        print("   row %-2d %-6s blue=%s expected_min=%s(%s) %s" % (
            r, ws.cell(row=r, column=1).value,
            [get_column_letter(x) for x in blue],
            get_column_letter(nm), nums[nm], "OK" if ok else "*** MISMATCH ***"))
    print("  mismatches:", bad)
