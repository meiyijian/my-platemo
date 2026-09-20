# -*- coding: utf-8 -*-
"""Verify the generated workbooks: layout, styles, values, blue cells."""
import os

import openpyxl

OUT_DIR = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本\PIEA最差版本与Nobatch最好版本"
SUFFIXES = ["最好", "均值"]

print("DIR contents:")
for n in sorted(os.listdir(OUT_DIR)):
    print("   %-70s %d bytes" % (n, os.path.getsize(os.path.join(OUT_DIR, n))))

for suffix in SUFFIXES:
    for m in (10, 15, 20):
        p = os.path.join(OUT_DIR, "IGDp%d目标_PIEA最差vsNoBatchDist%s.xlsx" % (m, suffix))
        print("=" * 120)
        print("FILE", os.path.basename(p), "exists", os.path.exists(p))
        if not os.path.exists(p):
            continue
        wb = openpyxl.load_workbook(p)
        ws = wb["IGDp"]
        print("  dims", ws.dimensions, "| merged", [str(x) for x in ws.merged_cells.ranges])
        print("  col widths", {k: round(v.width, 2) for k, v in ws.column_dimensions.items() if v.width})
        print("  row heights", sorted({v.height for v in ws.row_dimensions.values() if v.height}))
        blues = 0
        for r in range(1, ws.max_row + 1):
            line = []
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                col = cell.font.color
                rgb = col.rgb if (col is not None and col.type == "rgb") else None
                if rgb == "003333E9":
                    blues += 1
                line.append("%s%s" % (cell.value, "*" if rgb == "003333E9" else ""))
            print("  r%-3d" % r, " | ".join(str(x) for x in line))
        c1 = ws["D2"]
        print("  style check D2: font=%s/%s border(L=%s T=%s) align=%s/%s"
              % (c1.font.name, c1.font.sz, c1.border.left.style,
                 c1.border.top.style, c1.alignment.horizontal, c1.alignment.vertical))
        print("  blue cells: %d (expect exactly one per problem row = 16)" % blues)
