# -*- coding: utf-8 -*-
"""Compare the blue (row-minimum) marking of the source table and the new one."""
import os

import openpyxl

BASE = r"C:\Users\lsx\Desktop\AdaMao实验表\lambdat030版本"
PAIR = [("lambdat030nobatch十目标IGDp.xlsx", "lambdat030nobatch十目标IGDp_SAMOEA.xlsx")]

for src_name, new_name in PAIR:
    for name in (src_name, new_name):
        p = os.path.join(BASE, name)
        print("=" * 110)
        print(name)
        wb = openpyxl.load_workbook(p)
        ws = wb["IGDp"]
        for r in range(1, ws.max_row + 1):
            blues = []
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                col = cell.font.color
                rgb = col.rgb if (col is not None and col.type == "rgb") else None
                if rgb and rgb not in ("00000000", "FF000000"):
                    blues.append("%s(%s)" % (cell.column_letter, rgb))
            if blues:
                print("  r%-3d blue: %s" % (r, ", ".join(blues)))
        print("  fonts seen:", sorted({(ws.cell(row=r, column=c).font.name,
                                        ws.cell(row=r, column=c).font.sz)
                                       for r in (1, 2, 18)
                                       for c in range(1, ws.max_column + 1)}))
